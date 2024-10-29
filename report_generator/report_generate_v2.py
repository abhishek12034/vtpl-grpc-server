import os
import subprocess
from pylatex import Document, Section, NoEscape, Package, Command
import json
from pymongo import MongoClient
import cv2
import numpy as np
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
from docx.shared import Inches, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, ns
from docx.enum.table import WD_ALIGN_VERTICAL


class GenerateReport:
    # [Previous code remains exactly the same until the create_docx_report method's end]
    def __init__(self, config_file, db_descriptions):
        self.config_file = config_file 
        self.db_descriptions = db_descriptions      

    def load_db_data():
        # Connect to MongoDB
        client = MongoClient('mongodb://localhost:27017/')
        db = client['report']
        collection = db['report_db']

        # Retrieve all documents from the collection
        documents = collection.find()

        # Convert documents to a list of dictionaries
        data = [doc for doc in documents]

        return data[0] 

    def load_config(self):
        with open(self.config_file, 'r') as f:
            config = json.load(f)
            #todo:we have to sort the process names based on process index. future development
            process_names = [process['process_name'] for process in config['Processes']]
            print("Processes Name :", process_names)
            
        return process_names, config['Processes_meta']['input_output_image_show_report']

    def load_description_config(self):
        with open(self.db_descriptions, 'r') as f:
            descriptions = json.load(f)
            
        return descriptions

    def create_docx_report(self, output_dir, process_names, operations2, show_report):
        doc = DocxDocument()
        
        # Set up first section (title page) with no margins
        first_section = doc.sections[0]
        first_section.left_margin = Inches(0)
        first_section.right_margin = Inches(0)
        first_section.top_margin = Inches(0)
        first_section.bottom_margin = Inches(0)
        
        # Remove header/footer from first page
        first_section.different_first_page_header_footer = True
        first_header = first_section.header
        first_header.is_linked_to_previous = False
        first_header.paragraphs[0].text = ""
        first_footer = first_section.footer
        first_footer.is_linked_to_previous = False
        first_footer.paragraphs[0].text = ""
        
        # Add title page with full-page image
        title_paragraph = doc.add_paragraph()
        title_run = title_paragraph.add_run()
        title_image_path = 'D:/Projects/Forensic/vtpl-grpc-server/report_generator/images/Title_Page_Word.png'
        if os.path.exists(title_image_path):
            try:
                title_run.add_picture(title_image_path, width=Inches(8.5))
            except Exception as e:
                print(f"Error adding title image: {str(e)}")
                title = doc.add_heading('Image Adjustment Operations Report', 0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            print(f"Title image not found at path: {title_image_path}")
            title = doc.add_heading('Image Adjustment Operations Report', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add new section for TOC with normal margins
        toc_section = doc.add_section()
        toc_section.left_margin = Inches(1)
        toc_section.right_margin = Inches(1)
        toc_section.top_margin = Inches(1)
        toc_section.bottom_margin = Inches(1)
        toc_section.different_first_page_header_footer = False
        toc_section.start_page_number = 1  # Start page numbering from 1
        
        # Add header with page numbers starting from TOC
        toc_header = toc_section.header
        toc_header.is_linked_to_previous = False  # Unlink from previous section
        header_para = toc_header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header_para.add_run()
        
        # Start page numbering
        fld_char1 = OxmlElement('w:fldChar')
        fld_char1.set(qn('w:fldCharType'), 'begin')
        header_run._r.append(fld_char1)
        
        instr_text = OxmlElement('w:instrText')
        instr_text.set(qn('xml:space'), 'preserve')
        instr_text.text = 'PAGE'
        header_run._r.append(instr_text)
        
        fld_char2 = OxmlElement('w:fldChar')
        fld_char2.set(qn('w:fldCharType'), 'end')
        header_run._r.append(fld_char2)
        
        header_run = header_para.add_run(' of ')
        
        fld_char3 = OxmlElement('w:fldChar')
        fld_char3.set(qn('w:fldCharType'), 'begin')
        header_run._r.append(fld_char3)
        
        instr_text2 = OxmlElement('w:instrText')
        instr_text2.set(qn('xml:space'), 'preserve')
        instr_text2.text = 'NUMPAGES'
        header_run._r.append(instr_text2)
        
        fld_char4 = OxmlElement('w:fldChar')
        fld_char4.set(qn('w:fldCharType'), 'end')
        header_run._r.append(fld_char4)
        
        # Add footer with copyright text
        toc_footer = toc_section.footer
        toc_footer.is_linked_to_previous = False  # Unlink from previous section
        footer_para = toc_footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer_para.add_run('©2024-2025 Videonetics Technology Private Limited. All Rights Reserved')
        footer_run.font.size = Pt(8)
        footer_run.font.name = 'Arial'
        
        # Set up default document styles
        style = doc.styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor(0, 0, 0)  # Black color
        
        # Modify Heading 1 style
        heading1_style = doc.styles['Heading 1']
        heading1_style.font.name = 'Arial'
        heading1_style.font.size = Pt(14)
        heading1_style.font.bold = True
        heading1_style.font.color.rgb = RGBColor(0, 0, 255)  # Blue color
        
        # Add table of contents heading
        toc_heading = doc.add_heading('Table of Contents', level=1)
        toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add table of contents field
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        
        # Begin the TOC field
        fld_char = OxmlElement('w:fldChar')
        fld_char.set(qn('w:fldCharType'), 'begin')
        run._element.append(fld_char)
        
        # Add the TOC instruction text
        instr_text = OxmlElement('w:instrText')
        instr_text.set(qn('xml:space'), 'preserve')
        instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
        run._element.append(instr_text)
        
        # End the TOC field
        fld_char = OxmlElement('w:fldChar')
        fld_char.set(qn('w:fldCharType'), 'end')
        run._element.append(fld_char)
        
        # Add a section break for content
        content_section = doc.add_section()
        content_section.left_margin = Inches(1)
        content_section.right_margin = Inches(1)
        content_section.top_margin = Inches(1)
        content_section.bottom_margin = Inches(1)
        
        # Link header and footer to previous section (TOC)
        content_section.header.is_linked_to_previous = True
        content_section.footer.is_linked_to_previous = True
        
        # Calculate available width for images
        page_width = 8.5
        margin_width = 2
        available_width = page_width - margin_width
        image_width = Inches((available_width / 2) - 0.5)
        
        # Content sections
        for process_name in process_names:
            if process_name in operations2:
                operation = operations2[process_name]
                
                title = operation.get(process_name, operation['title'])
                description = operation.get(f"{process_name}_description", operation['description'])
                
                heading = doc.add_heading(title, level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                para = doc.add_paragraph(description)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.style = doc.styles['Normal']
                
                if show_report == "True":
                    table = doc.add_table(rows=2, cols=2)
                    table.style = 'Table Grid'
                    table.autofit = False
                    
                    # Set table properties to minimize spacing
                    table._element.xpath('.//w:tblLook')[0].set(ns.qn('w:firstRow'), '0')
                    table._element.xpath('.//w:tblLook')[0].set(ns.qn('w:lastRow'), '0')
                    
                    # Set minimal cell margins/padding
                    for row in table.rows:
                        for cell in row.cells:
                            cell.width = Inches(available_width / 2)
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            tc = cell._element
                            tcPr = tc.get_or_add_tcPr()
                            tcMar = OxmlElement('w:tcMar')
                            for side in ['top', 'right', 'bottom', 'left']:
                                node = OxmlElement(f'w:{side}')
                                node.set(ns.qn('w:w'), '50')
                                node.set(ns.qn('w:type'), 'dxa')
                                tcMar.append(node)
                            tcPr.append(tcMar)
                    
                    # Input image
                    input_cell = table.cell(0, 0)
                    input_para = input_cell.paragraphs[0]
                    input_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    input_para.space_before = Pt(0)
                    input_para.space_after = Pt(0)
                    
                    base_dir = os.path.abspath(os.path.dirname(__file__))
                    input_img_path = os.path.join(base_dir, 'report/channels_images', 'input_images', 
                                                f'input_{process_name.lower().replace(" ", "_")}.jpg')
                    
                    if os.path.exists(input_img_path):
                        try:
                            input_para.add_run().add_picture(input_img_path, width=image_width)
                            print(f"Successfully added input image: {input_img_path}")
                        except Exception as e:
                            print(f"Error adding input image for {process_name}: {str(e)}")
                            input_run = input_para.add_run("Error loading image")
                            input_run.font.name = 'Arial'
                            input_run.font.size = Pt(10)
                    else:
                        print(f"Input image not found at path: {input_img_path}")
                        input_run = input_para.add_run("Image not found")
                        input_run.font.name = 'Arial'
                        input_run.font.size = Pt(10)
                    
                    # Input label
                    input_label = table.cell(1, 0).paragraphs[0]
                    input_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    input_label.space_before = Pt(0)
                    input_label.space_after = Pt(0)
                    input_run = input_label.add_run("Input")
                    input_run.font.name = 'Arial'
                    input_run.font.size = Pt(10)
                    
                    # Output image
                    output_cell = table.cell(0, 1)
                    output_para = output_cell.paragraphs[0]
                    output_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    output_para.space_before = Pt(0)
                    output_para.space_after = Pt(0)
                    
                    output_img_path = os.path.join(base_dir, 'report/channels_images', 'output_images', 
                                                f'output_{process_name.lower().replace(" ", "_")}.jpg')
                    
                    if os.path.exists(output_img_path):
                        try:
                            output_para.add_run().add_picture(output_img_path, width=image_width)
                            print(f"Successfully added output image: {output_img_path}")
                        except Exception as e:
                            print(f"Error adding output image for {process_name}: {str(e)}")
                            output_run = output_para.add_run("Error loading image")
                            output_run.font.name = 'Arial'
                            output_run.font.size = Pt(10)
                    else:
                        print(f"Output image not found at path: {output_img_path}")
                        output_run = output_para.add_run("Image not found")
                        output_run.font.name = 'Arial'
                        output_run.font.size = Pt(10)
                    
                    # Output label
                    output_label = table.cell(1, 1).paragraphs[0]
                    output_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    output_label.space_before = Pt(0)
                    output_label.space_after = Pt(0)
                    output_run = output_label.add_run("Output")
                    output_run.font.name = 'Arial'
                    output_run.font.size = Pt(10)
                    
                    # Set minimal row heights
                    table.rows[0].height = Inches(0.5)
                    table.rows[1].height = Inches(0.1)
                
                # Add minimal spacing after each algorithm section
                spacing_para = doc.add_paragraph()
                spacing_para.space_after = Pt(12)
                spacing_para.style = doc.styles['Normal']
        
        # Add a section break for the last page
        last_section = doc.add_section()
        last_section.left_margin = Inches(0)
        last_section.right_margin = Inches(0)
        last_section.top_margin = Inches(0)
        last_section.bottom_margin = Inches(0)
        
        # Remove header/footer from last page
        last_section.different_first_page_header_footer = True
        last_header = last_section.header
        last_header.is_linked_to_previous = False
        last_header.paragraphs[0].text = ""
        last_footer = last_section.footer
        last_footer.is_linked_to_previous = False
        last_footer.paragraphs[0].text = ""
        
        # Add last page with full-page image
        last_page_paragraph = doc.add_paragraph()
        last_page_run = last_page_paragraph.add_run()
        last_image_path = 'D:/Projects/Forensic/vtpl-grpc-server/report_generator/images/Last_page.png'
        
        if os.path.exists(last_image_path):
            try:
                # Set both width and height for the last page image
                last_page_run.add_picture(last_image_path, width=Inches(8.5), height=Inches(11.0))
                print(f"Successfully added last page image: {last_image_path}")
            except Exception as e:
                print(f"Error adding last page image: {str(e)}")
                try:
                    # Fallback attempt with adjusted height
                    last_page_run.add_picture(last_image_path, width=Inches(8.5), height=Inches(7.5))
                    print("Successfully added last page image with adjusted dimensions")
                except Exception as e:
                    print(f"Error adding last page image with adjusted dimensions: {str(e)}")
        else:
            print(f"Last page image not found at path: {last_image_path}")
        
        # Save the document
        try:
            docx_path = os.path.join(output_dir, 'color_channels_report.docx')
            doc.save(docx_path)
            print(f"Word document generated successfully: {docx_path}")
            
            # Update the TOC
            word = None
            try:
                import win32com.client
                word = win32com.client.DispatchEx('Word.Application')
                doc = word.Documents.Open(docx_path)
                doc.TablesOfContents(1).Update()
                doc.Save()
                print("Table of Contents updated successfully")
            except Exception as e:
                print(f"Warning: Could not automatically update TOC: {str(e)}")
                print("Please open the document in Word and press F9 to update the Table of Contents")
            finally:
                if word:
                    doc.Close()
                    word.Quit()
                    
        except Exception as e:
            print(f"Error saving document: {str(e)}")

    def generate_report(self, output_dir='report'):
        output_dir = os.path.abspath(output_dir)
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        # Load configuration
        process_names, show_report = self.load_config()
        print("Show Report Flag :", show_report)
        operations2 = self.load_description_config()
        
        # Generate Word document
        self.create_docx_report(output_dir, process_names, operations2, show_report)
        
        geometry_options = {
            "margin": "2.54cm",
            "includeheadfoot": False
        }
        doc = Document(geometry_options=geometry_options)

        # [Previous package additions remain the same]
        # Add packages with enhanced styling
        doc.packages.append(Package('graphicx'))
        doc.packages.append(Package('float'))
        doc.packages.append(Package('xcolor'))
        doc.packages.append(Package('geometry'))
        doc.packages.append(Package('fancyhdr'))
        doc.packages.append(Package('titlesec'))
        doc.packages.append(Package('hyperref'))
        doc.packages.append(Package('helvet'))
        doc.packages.append(Package('lastpage'))  # Add this package for total page count

        # Enhanced document styling
        doc.preamble.append(NoEscape(r'\usepackage{xcolor}'))
        doc.preamble.append(NoEscape(r'\definecolor{headerblue}{RGB}{0,0,255}'))
        doc.preamble.append(NoEscape(r'\renewcommand{\familydefault}{\sfdefault}'))
        doc.preamble.append(NoEscape(r'\renewcommand{\sfdefault}{phv}'))
        
        # Setup fancy headers and footers
        doc.preamble.append(NoEscape(r'\pagestyle{fancy}'))
        doc.preamble.append(NoEscape(r'\fancyhf{}'))  # Clear all headers/footers
        doc.preamble.append(NoEscape(r'\fancyhead[R]{\textcolor{black!30}{page \thepage\ of \pageref{LastPage}}}'))
        doc.preamble.append(NoEscape(r'\fancyfoot[R]{\hspace{-1em}\textcolor{gray!20}{\small\copyright2024-2025 Videonetics Technology Private Limited. All Rights Reserved}}'))
        doc.preamble.append(NoEscape(r'\renewcommand{\headrulewidth}{0pt}'))
        doc.preamble.append(NoEscape(r'\renewcommand{\footrulewidth}{0pt}'))
        
        # No header/footer on title page and last page
        doc.preamble.append(NoEscape(r'\fancypagestyle{plain}{\fancyhf{}}'))

        # [Previous styling code remains the same]
        # Set up font sizes and styles
        doc.preamble.append(NoEscape(r'\usepackage{anyfontsize}'))
        doc.preamble.append(NoEscape(r'\usepackage{sectsty}'))
        # Set section headers to blue while keeping font settings
        doc.preamble.append(NoEscape(r'\sectionfont{\fontsize{14}{16}\selectfont\bfseries\color{headerblue}}'))
        # Set body text to black and 10pt
        doc.preamble.append(NoEscape(r'\fontsize{10}{12}\selectfont\color{black}'))
        
        # Custom section styling with blue color
        doc.preamble.append(NoEscape(r'\titleformat{\section}'
                                   r'{\color{headerblue}\fontsize{14}{16}\selectfont\bfseries\sffamily}'
                                   r'{\thesection}{1em}{}'))
        
        # Remove the empty page at the beginning
        doc.preamble.append(NoEscape(r'\AtBeginDocument{\let\cleardoublepage\clearpage}'))

        # Title page with the image
        doc.append(NoEscape(r'\begin{titlepage}'))
        doc.append(NoEscape(r'\thispagestyle{empty}'))
        doc.append(NoEscape(r'\newgeometry{margin=0pt}'))
        doc.append(NoEscape(r'\noindent\includegraphics[width=\paperwidth,height=\paperheight]{D:/Projects/Forensic/vtpl-grpc-server/report_generator/images/Title_Page_Pdf.png}'))
        doc.append(NoEscape(r'\restoregeometry'))
        doc.append(NoEscape(r'\end{titlepage}'))
        
        # Table of contents styling (also in blue for headers)
        doc.append(NoEscape(r'\hypersetup{linkcolor=headerblue}'))  # Make TOC entries blue
        doc.append(NoEscape(r'\tableofcontents'))
        doc.append(NoEscape(r'\clearpage'))

        def add_operation_section(title, description, input_img_path, output_img_path):
            with doc.create(Section(title)):
                doc.append('\n\n')
                doc.append(NoEscape(r'\begin{justify}'))
                doc.append(NoEscape(r'\large ' + description))
                doc.append(NoEscape(r'\end{justify}'))
                doc.append('\n\n')

                if show_report == "True":
                    print("INSIDE IF Condition :")
                    doc.append(NoEscape(r'\vspace{2em}'))
                    doc.append(NoEscape(r'\noindent\begin{minipage}{0.45\textwidth}'))
                    doc.append(NoEscape(r'\centering'))
                    doc.append(NoEscape(r'\includegraphics[width=\textwidth]{' + input_img_path + '}'))
                    doc.append(NoEscape(r'\\\fontsize{10}{12}\selectfont\color{black} Input'))  # Black caption  # 10pt for caption
                    doc.append(NoEscape(r'\end{minipage}'))
                    doc.append(NoEscape(r'\hfill'))
                    doc.append(NoEscape(r'\begin{minipage}{0.45\textwidth}'))
                    doc.append(NoEscape(r'\centering'))
                    doc.append(NoEscape(r'\includegraphics[width=\textwidth]{' + output_img_path + '}'))
                    doc.append(NoEscape(r'\\\fontsize{10}{12}\selectfont\color{black} Output'))  # Black caption  # 10pt for caption
                    doc.append(NoEscape(r'\end{minipage}'))
                    doc.append(NoEscape(r'\vspace{1em}'))

        # Add each operation to the document
        for process_name in process_names:
            if process_name in operations2:
                operation = operations2[process_name]
                
                print(f"Processing: {process_name}")

                title = operation.get(process_name, operation['title'])
                description = operation.get(f"{process_name}_description", operation['description'])
                input_image_path = f'./channels_images/input_images/input_{process_name.lower().replace(" ", "_")}.jpg'
                output_image_path = f'./channels_images/output_images/output_{process_name.lower().replace(" ", "_")}.jpg'
                print("LATEX Output IMG Path :", output_image_path)
                
                add_operation_section(
                    title,
                    description,
                    input_image_path,
                    output_image_path
                )

        # Add last page with full-page image
        doc.append(NoEscape(r'\clearpage'))  # Ensure new page
        doc.append(NoEscape(r'\thispagestyle{empty}'))  # Remove header/footer
        doc.append(NoEscape(r'\newgeometry{margin=0pt}'))  # Remove margins
        doc.append(NoEscape(r'\noindent\includegraphics[width=\paperwidth,height=\paperheight]{D:/Projects/Forensic/vtpl-grpc-server/report_generator/images/Last_page.png}'))
        doc.append(NoEscape(r'\restoregeometry'))

        # Generate TEX file
        tex_filename = 'color_channels_report1.tex'
        tex_fullpath = os.path.join(output_dir, tex_filename)

        # Save the TEX file
        doc.generate_tex(os.path.splitext(tex_fullpath)[0])

        # Compile the TEX file to PDF
        pdflatex_path = r"C:\Users\katan\AppData\Local\Programs\MiKTeX\miktex\bin\x64\pdflatex.exe"
        try:
            for _ in range(3):  # Run 3 times to ensure TOC is properly generated
                process = subprocess.run(
                    [pdflatex_path, 
                    '-interaction=nonstopmode',
                    tex_filename],
                    check=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    cwd=output_dir,
                    text=True
                )
                print(process.stdout)

            pdf_path = os.path.join(output_dir, 'color_channels_report.pdf')
            if os.path.exists(pdf_path):
                print(f"PDF generated successfully: {pdf_path}")
            else:
                print(f"PDF file not found at expected location: {pdf_path}")

        except subprocess.CalledProcessError as e:
            print(f"Error compiling PDF: {e}")
            print(f"stdout: {e.stdout}")
            print(f"stderr: {e.stderr}")
            print(f"Current working directory: {os.getcwd()}")
            print(f"Output directory: {output_dir}")
            print(f"TEX file exists: {os.path.exists(tex_fullpath)}")

if __name__ == "__main__":
    report_obj = GenerateReport(config_file='./config.json', db_descriptions='./new_descriptions.json')
    report_obj.generate_report()